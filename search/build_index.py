"""Build an Azure AI Search evidence index over a local data-room folder.

Walks a folder, extracts text per page (PDF via pypdf, DOCX via python-docx, XLSX via openpyxl,
MD/TXT directly; scanned PDFs and images optionally via Document Intelligence prebuilt-read),
chunks it with overlap while keeping the page number, and pushes the chunks to an Azure AI Search
index over REST. Every chunk carries source_path + page so the Foundry agent can cite clause/page.

Search REST api-version 2026-04-01 (GA):
  https://learn.microsoft.com/en-us/rest/api/searchservice/search-service-api-versions
  PUT  {endpoint}/indexes('{name}')?api-version=...            create/update schema
  POST {endpoint}/indexes('{name}')/docs/search.index?api-version=...  mergeOrUpload chunks

Auth: --api-key reads AZURE_SEARCH_API_KEY (never printed); otherwise Entra ID token for
https://search.azure.com/.default via DefaultAzureCredential (--auth default) or
InteractiveBrowserCredential (--auth interactive, optional --client-id / --tenant-id).

Examples:
  # offline: extract + chunk to JSONL, no Azure calls
  uv run --no-project --with pypdf --with python-docx --with openpyxl \
      python search/build_index.py --root ./dataroom --dry-run --out out/search
  # upload
  uv run --no-project --with pypdf --with python-docx --with openpyxl --with azure-identity \
      python search/build_index.py --root ./dataroom --endpoint https://<svc>.search.windows.net \
      --index spine-evidence --auth interactive
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Dict, Iterable, Iterator, List, Optional, Sequence, Tuple

if __package__ in (None, ""):  # allow `python search/build_index.py`
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from ingest.azure_common import LONG_PATH_THRESHOLD, HeaderFn, bearer_headers, http_json, key_headers, long_path  # noqa: E402

SEARCH_API_VERSION = "2026-04-01"
SEARCH_SCOPE = "https://search.azure.com/.default"
SEARCH_KEY_ENV = "AZURE_SEARCH_API_KEY"
SEMANTIC_CONFIG = "evidence-semantic"

TEXT_EXT = {".md", ".txt", ".csv"}
PDF_EXT = {".pdf"}
DOCX_EXT = {".docx"}
XLSX_EXT = {".xlsx", ".xlsm"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
SUPPORTED_EXT = TEXT_EXT | PDF_EXT | DOCX_EXT | XLSX_EXT | IMAGE_EXT

Pages = List[Tuple[int, str]]


@dataclass
class Chunk:
    id: str
    content: str
    source_path: str
    file_name: str
    page: int
    doc_type: str
    folder: str
    sha256: str
    chunk_no: int
    last_modified: str


# ---------------------------------------------------------------------------------------------
# chunking (pure, tested offline)
# ---------------------------------------------------------------------------------------------


def normalise_ws(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def split_text(text: str, size: int = 1200, overlap: int = 200, min_size: Optional[int] = None) -> List[str]:
    """Split text into windows of ~size chars with `overlap` chars shared between neighbours.

    Breaks at a paragraph, line, sentence or word boundary found in the last part of the
    window, so chunks usually land between min_size (default 0.8*size) and size characters.
    """
    if size <= 0 or overlap < 0 or overlap >= size:
        raise ValueError("require size > 0 and 0 <= overlap < size")
    text = normalise_ws(text)
    if not text:
        return []
    min_size = int(size * 0.8) if min_size is None else min_size
    out: List[str] = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + size, n)
        if end < n:
            window = text[start:end]
            cut = -1
            for sep in ("\n\n", "\n", ". ", "; ", " "):
                i = window.rfind(sep, min_size)
                if i != -1:
                    cut = i + len(sep)
                    break
            if cut > 0:
                end = start + cut
        piece = text[start:end].strip()
        if piece:
            out.append(piece)
        if end >= n:
            break
        nxt = end - overlap
        # snap the overlap start forward to a word boundary so chunks do not start mid-word
        sp = text.find(" ", nxt, end)
        start = sp + 1 if sp != -1 else max(nxt, start + 1)
    return out


def doc_id(source_path: str, chunk_no: int) -> str:
    """Stable, key-safe (letters, digits, '_', '-') id per (path, chunk)."""
    h = hashlib.sha1(source_path.replace("\\", "/").lower().encode("utf-8")).hexdigest()[:24]
    return f"{h}-{chunk_no:05d}"


def chunk_pages(
    pages: Pages,
    *,
    source_path: str,
    folder: str,
    doc_type: str,
    sha256: str,
    last_modified: str,
    size: int = 1200,
    overlap: int = 200,
) -> List[Chunk]:
    """Chunk each page separately so every chunk maps to exactly one page number."""
    chunks: List[Chunk] = []
    file_name = source_path.replace("\\", "/").rsplit("/", 1)[-1]
    for page, text in pages:
        for piece in split_text(text, size=size, overlap=overlap):
            n = len(chunks)
            chunks.append(
                Chunk(
                    id=doc_id(source_path, n),
                    content=piece,
                    source_path=source_path,
                    file_name=file_name,
                    page=int(page),
                    doc_type=doc_type,
                    folder=folder,
                    sha256=sha256,
                    chunk_no=n,
                    last_modified=last_modified,
                )
            )
    return chunks


def index_schema(name: str) -> Dict[str, object]:
    """Index definition: keyword + semantic ranking over `content`, filters on provenance fields."""

    def f(name_: str, type_: str, **kw: object) -> Dict[str, object]:
        base: Dict[str, object] = {"name": name_, "type": type_, "retrievable": True, "searchable": False, "filterable": False, "sortable": False, "facetable": False}
        base.update(kw)
        return base

    return {
        "name": name,
        "fields": [
            f("id", "Edm.String", key=True, filterable=True),
            f("content", "Edm.String", searchable=True, analyzer="en.microsoft"),
            f("source_path", "Edm.String", filterable=True),
            f("file_name", "Edm.String", searchable=True, filterable=True, sortable=True),
            f("page", "Edm.Int32", filterable=True, sortable=True),
            f("doc_type", "Edm.String", filterable=True, facetable=True),
            f("folder", "Edm.String", searchable=True, filterable=True, facetable=True),
            f("sha256", "Edm.String", filterable=True),
            f("chunk_no", "Edm.Int32", filterable=True, sortable=True),
            f("last_modified", "Edm.DateTimeOffset", filterable=True, sortable=True),
        ],
        "semantic": {
            "defaultConfiguration": SEMANTIC_CONFIG,
            "configurations": [
                {
                    "name": SEMANTIC_CONFIG,
                    "prioritizedFields": {
                        "titleField": {"fieldName": "file_name"},
                        "prioritizedContentFields": [{"fieldName": "content"}],
                        "prioritizedKeywordsFields": [{"fieldName": "folder"}, {"fieldName": "doc_type"}],
                    },
                }
            ],
        },
    }


# ---------------------------------------------------------------------------------------------
# file walking and extraction
# ---------------------------------------------------------------------------------------------


def doc_type_of(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in PDF_EXT:
        return "pdf"
    if ext in DOCX_EXT:
        return "docx"
    if ext in XLSX_EXT:
        return "xlsx"
    if ext in IMAGE_EXT:
        return "image"
    return ext.lstrip(".") or "unknown"


def _strip_long_prefix(s: str) -> str:
    if s.startswith("\\\\?\\UNC\\"):
        return "\\\\" + s[8:]
    return s.removeprefix("\\\\?\\")


def iter_files(root: Path, on_error: Optional[Callable[[OSError], None]] = None) -> Iterator[Path]:
    """Walk root (always through the \\\\?\\ prefix on Windows so folders past MAX_PATH are visited)."""
    for dirpath, dirnames, filenames in os.walk(long_path(root, force=True), onerror=on_error):
        dirnames[:] = sorted(d for d in dirnames if not d.startswith((".", "~$")))
        for fn in sorted(filenames):
            if fn.startswith(("~$", ".")):
                continue
            yield Path(_strip_long_prefix(dirpath)) / fn


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(long_path(path), "rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


class ShortPath:
    """Context manager yielding a path that third-party parsers can open (temp copy when too long)."""

    def __init__(self, path: Path) -> None:
        self.path = path
        self._tmp: Optional[str] = None

    def __enter__(self) -> str:
        if os.name == "nt" and len(os.path.abspath(str(self.path))) > LONG_PATH_THRESHOLD:
            self._tmp = tempfile.mkdtemp(prefix="idx_")
            dst = os.path.join(self._tmp, "f" + self.path.suffix.lower())
            shutil.copyfile(long_path(self.path), dst)
            return dst
        return str(self.path)

    def __exit__(self, *exc: object) -> None:
        if self._tmp:
            shutil.rmtree(self._tmp, ignore_errors=True)


def extract_pdf(path: str) -> Pages:
    from pypdf import PdfReader  # lazy: optional extra

    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as e:  # noqa: BLE001
            raise ValueError(f"encrypted PDF: {e}") from None
    return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]


def extract_docx(path: str) -> Pages:
    """DOCX has no fixed pagination: explicit or last-rendered page breaks advance the counter.

    Tables are appended to the last page (python-docx does not expose body order cheaply).
    """
    import docx  # python-docx, lazy

    d = docx.Document(path)
    pages: Pages = []
    page, buf = 1, []
    for para in d.paragraphs:
        xml = para._p.xml
        if buf and ('w:type="page"' in xml or "w:lastRenderedPageBreak" in xml):
            pages.append((page, "\n".join(buf)))
            page, buf = page + 1, []
        if para.text.strip():
            buf.append(para.text)
    for t_i, table in enumerate(d.tables):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        buf.append(f"[table {t_i + 1}]\n" + "\n".join(rows))
    if buf:
        pages.append((page, "\n".join(buf)))
    return pages


def extract_xlsx(path: str, max_rows: int = 5000) -> Pages:
    """One 'page' per sheet (page = sheet index, 1-based); rows rendered as 'a | b | c'."""
    from openpyxl import load_workbook  # lazy

    wb = load_workbook(path, read_only=True, data_only=True)
    pages: Pages = []
    try:
        for s_i, ws in enumerate(wb.worksheets):
            lines = [f"[sheet {ws.title}]"]
            for r_i, row in enumerate(ws.iter_rows(values_only=True)):
                if r_i >= max_rows:
                    lines.append(f"... truncated after {max_rows} rows")
                    break
                cells = ["" if v is None else str(v) for v in row]
                if any(c.strip() for c in cells):
                    lines.append(" | ".join(cells).rstrip(" |"))
            pages.append((s_i + 1, "\n".join(lines)))
    finally:
        wb.close()
    return pages


def extract_text_file(path: str) -> Pages:
    with open(path, "rb") as fh:
        raw = fh.read()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return [(1, raw.decode(enc))]
        except UnicodeDecodeError:
            continue
    return [(1, raw.decode("utf-8", "replace"))]


def extract_pages(path: Path, ocr: Optional[Callable[[bytes], Pages]] = None, min_chars_per_page: int = 25) -> Tuple[Pages, str]:
    """Return (pages, method). Falls back to OCR for scanned PDFs / images when `ocr` is given."""
    ext = path.suffix.lower()
    with ShortPath(path) as p:
        if ext in PDF_EXT:
            pages = extract_pdf(p)
            if ocr and (not pages or sum(len(t.strip()) for _, t in pages) < min_chars_per_page * max(len(pages), 1)):
                with open(p, "rb") as fh:
                    return ocr(fh.read()), "docintel-read"
            return pages, "pypdf"
        if ext in DOCX_EXT:
            return extract_docx(p), "python-docx"
        if ext in XLSX_EXT:
            return extract_xlsx(p), "openpyxl"
        if ext in TEXT_EXT:
            return extract_text_file(p), "text"
        if ext in IMAGE_EXT:
            if not ocr:
                raise ValueError("image requires --docintel-endpoint for OCR")
            with open(p, "rb") as fh:
                return ocr(fh.read()), "docintel-read"
    raise ValueError(f"unsupported extension {ext}")


# ---------------------------------------------------------------------------------------------
# Azure AI Search REST
# ---------------------------------------------------------------------------------------------


def ensure_index(endpoint: str, name: str, headers: HeaderFn, api_version: str = SEARCH_API_VERSION) -> None:
    url = f"{endpoint.rstrip('/')}/indexes('{name}')?api-version={api_version}"
    http_json("PUT", url, headers(), body=index_schema(name))


def upload(endpoint: str, name: str, chunks: Sequence[Chunk], headers: HeaderFn, batch: int = 500, api_version: str = SEARCH_API_VERSION) -> Tuple[int, List[str]]:
    url = f"{endpoint.rstrip('/')}/indexes('{name}')/docs/search.index?api-version={api_version}"
    ok, errors = 0, []
    for i in range(0, len(chunks), batch):
        payload = {"value": [{"@search.action": "mergeOrUpload", **asdict(c)} for c in chunks[i : i + batch]]}
        _, _, body = http_json("POST", url, headers(), body=payload)
        for r in (body or {}).get("value", []):
            if r.get("status"):
                ok += 1
            else:
                errors.append(f"{r.get('key')}: {r.get('statusCode')} {r.get('errorMessage')}")
    return ok, errors


# ---------------------------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------------------------


def build_chunks(
    root: Path,
    *,
    max_mb: float,
    size: int,
    overlap: int,
    ocr: Optional[Callable[[bytes], Pages]],
    report: List[Dict[str, str]],
) -> Iterable[Chunk]:
    root = root.resolve()
    def walk_error(e: OSError) -> None:
        report.append({"source_path": str(getattr(e, "filename", "")), "status": "unreadable", "reason": f"cannot list folder: {e.strerror}", "size_bytes": ""})

    for path in iter_files(root, on_error=walk_error):
        rel = os.path.relpath(str(path), str(root)).replace("\\", "/")
        ext = path.suffix.lower()
        try:
            st = os.stat(long_path(path))
        except OSError as e:
            report.append({"source_path": rel, "status": "unreadable", "reason": f"stat failed: {e}", "size_bytes": ""})
            continue
        if ext not in SUPPORTED_EXT:
            report.append({"source_path": rel, "status": "skipped", "reason": f"unsupported type {ext or '(none)'}", "size_bytes": str(st.st_size)})
            continue
        if st.st_size > max_mb * 1024 * 1024:
            report.append({"source_path": rel, "status": "skipped", "reason": f"larger than {max_mb} MB", "size_bytes": str(st.st_size)})
            continue
        if ext in IMAGE_EXT and not ocr:
            report.append({"source_path": rel, "status": "skipped", "reason": "image without --docintel-endpoint", "size_bytes": str(st.st_size)})
            continue
        try:
            pages, method = extract_pages(path, ocr=ocr)
        except Exception as e:  # noqa: BLE001 - any parser failure is logged, never fatal
            report.append({"source_path": rel, "status": "unreadable", "reason": f"{type(e).__name__}: {e}"[:500], "size_bytes": str(st.st_size)})
            continue
        chunks = chunk_pages(
            pages,
            source_path=rel,
            folder=rel.rsplit("/", 1)[0] if "/" in rel else "",
            doc_type=doc_type_of(path),
            sha256=sha256_file(path),
            last_modified=datetime.fromtimestamp(st.st_mtime, tz=timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
            size=size,
            overlap=overlap,
        )
        if not chunks:
            report.append({"source_path": rel, "status": "no_text", "reason": f"no extractable text via {method} (scanned? pass --docintel-endpoint)", "size_bytes": str(st.st_size)})
            continue
        report.append({"source_path": rel, "status": "indexed", "reason": f"{method}; {len(pages)} pages; {len(chunks)} chunks", "size_bytes": str(st.st_size)})
        yield from chunks


def search_headers(a: argparse.Namespace) -> HeaderFn:
    if a.api_key:
        h = key_headers("api-key", SEARCH_KEY_ENV)
        if h is None:
            raise SystemExit(f"--api-key given but {SEARCH_KEY_ENV} is not set")
        return h
    return bearer_headers(a.auth, SEARCH_SCOPE, a.client_id, a.tenant_id)


def main(argv: Optional[Sequence[str]] = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--root", type=Path, required=True, help="data-room folder to walk")
    ap.add_argument("--out", type=Path, default=Path("out/search"), help="folder for the report / dry-run JSONL")
    ap.add_argument("--endpoint", default=os.environ.get("AZURE_SEARCH_ENDPOINT", ""), help="https://<svc>.search.windows.net")
    ap.add_argument("--index", default=os.environ.get("AI_SEARCH_INDEX_NAME", "spine-evidence"))
    ap.add_argument("--auth", choices=["default", "interactive"], default="default")
    ap.add_argument("--client-id", default=None, help="app (interactive) or managed identity (default) client id")
    ap.add_argument("--tenant-id", default=None)
    ap.add_argument("--api-key", action="store_true", help=f"use admin key from env {SEARCH_KEY_ENV} instead of Entra ID")
    ap.add_argument("--api-version", default=SEARCH_API_VERSION)
    ap.add_argument("--docintel-endpoint", default="", help="enable OCR of scanned PDFs/images via prebuilt-read")
    ap.add_argument("--max-mb", type=float, default=50.0, help="skip files larger than this")
    ap.add_argument("--chunk-size", type=int, default=1200)
    ap.add_argument("--overlap", type=int, default=200)
    ap.add_argument("--dry-run", action="store_true", help="write chunks to <out>/chunks.jsonl; no Azure calls")
    ap.add_argument("--no-schema", action="store_true", help="do not create/update the index schema")
    a = ap.parse_args(argv)

    if not a.dry_run and not a.endpoint:
        ap.error("--endpoint (or AZURE_SEARCH_ENDPOINT) is required unless --dry-run")
    if not os.path.isdir(long_path(a.root)):
        ap.error(f"--root {a.root} is not a folder")

    ocr = None
    if a.docintel_endpoint:
        from ingest.docintel import analyze, di_headers, page_texts

        di_h = di_headers(a.auth, a.client_id, a.tenant_id)
        ocr = lambda data: page_texts(analyze(a.docintel_endpoint, "prebuilt-read", data, di_h))  # noqa: E731

    a.out.mkdir(parents=True, exist_ok=True)
    report: List[Dict[str, str]] = []
    chunks = list(build_chunks(a.root, max_mb=a.max_mb, size=a.chunk_size, overlap=a.overlap, ocr=ocr, report=report))

    report_path = a.out / "index_report.csv"
    with open(report_path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["source_path", "status", "reason", "size_bytes"])
        w.writeheader()
        w.writerows(report)
    counts: Dict[str, int] = {}
    for r in report:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    print(f"files: {counts}; chunks: {len(chunks)}; report: {report_path}")

    if a.dry_run:
        jsonl = a.out / "chunks.jsonl"
        with open(jsonl, "w", encoding="utf-8") as fh:
            for c in chunks:
                fh.write(json.dumps(asdict(c), ensure_ascii=False) + "\n")
        print(f"dry run: wrote {jsonl}")
        return 0

    headers = search_headers(a)
    if not a.no_schema:
        ensure_index(a.endpoint, a.index, headers, a.api_version)
        print(f"index '{a.index}' created/updated (api-version {a.api_version})")
    ok, errors = upload(a.endpoint, a.index, chunks, headers, api_version=a.api_version)
    print(f"uploaded {ok}/{len(chunks)} chunks")
    for e in errors[:20]:
        print(f"  failed: {e}", file=sys.stderr)
    return 0 if not errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
